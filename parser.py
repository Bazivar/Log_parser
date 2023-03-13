import re
import os
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.writer.excel import save_workbook
import webbrowser
import openpyxl.styles.numbers

from datetime import datetime


content = os.listdir()
for element in content:
    if element[-3:] == 'log':
        address = element

        f = open(element, 'r')

        file = f.read()
        f.close()

        date = re.findall(r'DATE_TIME_STAMP\[7\]=(\d*-\d*-\d*\s\d*:\d*:\d*)', file)
        voltage_VA_N = re.findall(r'VOLTAGE_VA_N\[4\]=(\d*[\.]*\d*)' ,file)
        voltage_VB_N = re.findall(r'VOLTAGE_VB_N\[4\]=(\d*[\.]*\d*)' ,file)
        voltage_VC_N = re.findall(r'VOLTAGE_VC_N\[4\]=(\d*[\.]*\d*)' ,file)
        voltage_VA_B = re.findall(r'VOLTAGE_VA_B\[4\]=(\d*[\.]*\d*)' ,file)
        voltage_VB_C = re.findall(r'VOLTAGE_VB_C\[4\]=(\d*[\.]*\d*)' ,file)
        voltage_VC_A = re.findall(r'VOLTAGE_VC_A\[4\]=(\d*[\.]*\d*)' ,file)
        frequency = re.findall(r'FREQUENCY\[4\]=(\d*[\.]*\d*)' ,file)


        FILE_NAME = address + '.xlsx'

        #creating an Excel file

        wb = openpyxl.load_workbook(filename='template.xlsx')

        data = wb['Данные']
        life_neutral = wb['Фаза-Нейтраль']
        life_life = wb['Фаза-Фаза']
        freq = wb['Частота']

        #making the data table
        data['A1'] = 'Дата и время'
        data['B1'] = 'Напряжение A-N'
        data['C1'] = 'Напряжение B-N'
        data['D1'] = 'Напряжение C-N'
        data['E1'] = 'Напряжение A-B'
        data['F1'] = 'Напряжение B-C'
        data['G1'] = 'Напряжение C-A'
        data['H1'] = 'Частота'
        for i in range (1, len(date)):
            data[f'A{i+1}'] = datetime_object = datetime.strptime(date[i-1], '%Y-%m-%d %H:%M:%S')
            data[f'B{i+1}'] = float(voltage_VA_N[i-1])
            data[f'C{i+1}'] = float(voltage_VB_N[i-1])
            data[f'D{i+1}'] = float(voltage_VC_N[i-1])
            data[f'E{i+1}'] = float(voltage_VA_B[i-1])
            data[f'F{i+1}'] = float(voltage_VB_C[i-1])
            data[f'G{i+1}'] = float(voltage_VC_A[i-1])
            data[f'H{i+1}'] = float(frequency[i-1])


        #creating the Line-Neutral chart
        chart = ScatterChart()
        chart.title = "Напряжение Фаза-Нейтраль "
        chart.style = 10
        chart.width = 40
        chart.height = 20
        chart.x_axis.title = "Дата и время"
        chart.y_axis.title = "Напряжение, В"
        chart.x_axis.number_format = 'd-mmm-h:mm:ss'

        xvalues = Reference(data, min_col=1, min_row=2, max_row=len(date))
        for i in range(2, 5):
            values = Reference(data, min_col=i, min_row=1, max_row=len(date))
            series = Series(values, xvalues, title_from_data=True)
            chart.series.append(series)
        life_neutral.add_chart(chart, 'C2')



        #creating the Life-Life chart
        chart = ScatterChart()
        chart.title = "Напряжение Фаза-Фаза "
        chart.style = 10
        chart.width = 40
        chart.height = 20
        chart.x_axis.title = "Дата и время"
        chart.y_axis.title = "Напряжение, В"
        chart.x_axis.number_format = 'd-mmm-h:mm:ss'

        xvalues = Reference(data, min_col=1, min_row=2, max_row=len(date))
        for i in range(5, 8):
            values = Reference(data, min_col=i, min_row=1, max_row=len(date))
            series = Series(values, xvalues, title_from_data=True)
            chart.series.append(series)
        life_life.add_chart(chart, 'C2')


        #creating the frequency chart
        chart = ScatterChart()
        chart.title = "Частота"
        chart.style = 10
        chart.width = 40
        chart.height = 20
        chart.x_axis.title = "Дата и время"
        chart.y_axis.title = "Частота, Гц"
        chart.x_axis.number_format = 'd-mmm-h:mm:ss'

        xvalues = Reference(data, min_col=1, min_row=2, max_row=len(date))
        for i in range(8, 9):
            values = Reference(data, min_col=i, min_row=1, max_row=len(date))
            series = Series(values, xvalues, title_from_data=True)
            chart.series.append(series)
        freq.add_chart(chart, 'C2')


        #saving the Excel file
        save_workbook(wb, FILE_NAME)

        #opening the Excel file on the PC
        webbrowser.open(FILE_NAME)
